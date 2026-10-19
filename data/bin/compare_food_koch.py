"""Compare our food decomposition (`ours`, Agribalyse 3.2 through Brightway) with the
ecolysis workbook of Peter Koch (`koch`, Agribalyse 4.0 through volca).

    uv run python data/bin/decompose_food.py decompose --all-food
    uv run python data/bin/compare_food_koch.py [--koch-workbook <path>]

Products are joined on their exact activity name. The workbook's `optionA` sheet gives
the totals per kg of product (electricity kWh, heat MJ, tap water kg), its
`detailedResults` sheet the inputs of each transformation step, from which Koch's main
raw activity and main ingredient are computed as for ours (`ingredient_key`). The Excel
written has two sheets:

- `ours_vs_koch`: one row per product, with a status: match (main ingredient kg,
  electricity, heat and water all within the tolerance), delta, ours only or koch only,
  the quantities that are off, the SimaPro category, and the main raw activity and main
  ingredient of both sides
- `inputs`: for every Koch product, one row per input class and input activity, aligned
  with ours on the exact name. Water, biowaste and transport are compared on their
  totals, as our decomposition keeps no detail for them.

The category is read from Brightway, so `BRIGHTWAY2_DIR` must be set as for
`decompose_food.py`. The workbook received from Koch is versioned in `data/food/koch/`.
Run from the repo root. Reads and writes in the gitignored `output/decompose_food/`.
"""

import json
import math
from pathlib import Path

import bw2data
import openpyxl
import typer

from config import DATA_ROOT_DIR, settings
from ecobalyse_data.food_decomposition import BrightwayGraph, Totals

OUTPUT_DIR = DATA_ROOT_DIR.parent / "output" / "decompose_food"
KOCH_WORKBOOK = DATA_ROOT_DIR / "food/koch/194-p09-ResultatsAnalyse10092026.xlsx"

QUANTITIES = (
    "main_raw_activity_kg",
    "main_ingredient_kg",
    "elec_kwh",
    "heat_mj",
    "water_kg",
)
# The status uses the main ingredient, the quantity components use. The main raw activity
# is shown to explain a gap (one origin vs all the origins of the ingredient).
STATUS_QUANTITIES = ("main_ingredient_kg", "elec_kwh", "heat_mj", "water_kg")
SIGNIFICANT_DIGITS = 4

# optionA: B product, E electricity kWh, F heat MJ, G tap water kg
TOTALS_SHEET, TOTALS_FIRST_ROW, TOTALS_NAME_COLUMN = "optionA", 4, 1
TOTALS_COLUMNS = {"elec_kwh": 4, "heat_mj": 5, "water_kg": 6}
# detailedResults: B product, D input type, E input activity, F amount
DETAIL_SHEET, DETAIL_FIRST_ROW = "detailedResults", 2
DETAIL_PRODUCT_COLUMN, DETAIL_TYPE_COLUMN = 1, 3
DETAIL_INPUT_COLUMN, DETAIL_AMOUNT_COLUMN = 4, 5
# Koch's input types and our classes. His other types (biosphereflow, coproducts, Input…)
# are not inputs of the decomposition.
KOCH_INPUT_CLASSES = {
    "AgricFoodRaw": "raw",
    "EnergyElectr": "elec",
    "EnergyHeat": "heat",
    "WaterRaw": "water",
    "BiowasteOrBy": "biowaste",
    "OtherTranspo": "transport",
    "OtherMateria": "other",
    "Other-divers": "other",
}
# detailedResults keeps electricity in MJ and transport in kg × m
KOCH_UNIT_FACTORS = {"elec": 1 / 3.6, "transport": 1 / 1_000_000}

# Classes compared input by input, and the key of our input detail
DETAILED_CLASSES = {"raw": "raw", "elec": "elec_detail", "heat": "heat_detail"}
# Classes compared on their total only, and the key of our total
TOTAL_ONLY_CLASSES = {
    "water": "water_kg",
    "biowaste": "biowaste_kg",
    "transport": "transport_tkm",
}
CLASS_ORDER = ["raw", "elec", "heat", "water", "biowaste", "transport", "other"]
TOTAL_INPUT_NAME = "(total)"

STATUSES = ("delta", "match", "ours only", "koch only")
COLUMNS = [
    "status",
    "delta_quantities",
    "ours_name",
    "koch_name",
    "ours_alias",
    "category",
    "ours_unit",
    *(
        f"{side}_{quantity}"
        for quantity in QUANTITIES
        for side in ("ours", "koch", "delta_pct")
    ),
    "ours_main_raw_activity",
    "koch_main_raw_activity",
    "ours_main_ingredient",
    "koch_main_ingredient",
    "same_main_ingredient",
    "ours_main_ingredient_share",
    "ours_raw_count",
    "ours_truncated",
]
INPUT_COLUMNS = [
    "ours_name",
    "koch_name",
    "class",
    "input_name",
    "ours_amount",
    "koch_amount",
    "delta_pct",
    "side",
]


def read_koch(total_rows, detail_rows) -> dict[str, dict]:
    """Koch's products by activity name, from the rows (tuples) of `optionA` and
    `detailedResults`. Each has its totals and its inputs by (class, name) with their
    amount in our units."""
    koch_products: dict[str, dict] = {}
    for row in total_rows:
        name = row[TOTALS_NAME_COLUMN]
        if name:
            totals = {
                quantity: float(row[column] or 0)
                for quantity, column in TOTALS_COLUMNS.items()
            }
            koch_products[name] = {**totals, "name": name, "inputs": {}}
    for row in detail_rows:
        product = koch_products.get(row[DETAIL_PRODUCT_COLUMN])
        if product is None:
            continue
        input_class = KOCH_INPUT_CLASSES.get(row[DETAIL_TYPE_COLUMN])
        if input_class is None:
            continue
        amount = float(row[DETAIL_AMOUNT_COLUMN] or 0)
        amount *= KOCH_UNIT_FACTORS.get(input_class, 1)
        key = (input_class, row[DETAIL_INPUT_COLUMN])
        product["inputs"][key] = product["inputs"].get(key, 0.0) + amount
    return koch_products


def koch_raw(koch: dict) -> dict[str, float]:
    return {
        name: amount
        for (input_class, name), amount in koch["inputs"].items()
        if input_class == "raw"
    }


def significant(value: float | None) -> float | None:
    """Round to 4 significant figures, to keep float noise out of the sheet."""
    if not value or not math.isfinite(value):
        return value
    return round(value, SIGNIFICANT_DIGITS - 1 - math.floor(math.log10(abs(value))))


def delta_pct(ours: float, koch: float) -> float | None:
    """Delta of ours against koch in percent, None when koch is 0 and ours is not."""
    if koch == 0:
        return 0.0 if ours == 0 else None
    return round((ours - koch) / koch * 100, 1)


def main_raws(raw: dict[str, float]) -> dict:
    """The main raw activity and main ingredient of raw inputs by name, computed by
    `Totals` as for the components, so both sides follow the same rule."""
    totals = Totals(raw=raw)
    main_raw_activity, main_raw_activity_kg = totals.main_raw_activity or (None, 0.0)
    main_ingredient, main_ingredient_kg = totals.main_ingredient or (None, 0.0)
    return {
        "main_raw_activity": main_raw_activity,
        "main_raw_activity_kg": main_raw_activity_kg,
        "main_ingredient": main_ingredient,
        "main_ingredient_kg": main_ingredient_kg,
    }


def category_name(category_path: tuple[str, ...]) -> str:
    """Third level of the SimaPro path (`Agricultural / Food / <category>`)."""
    if len(category_path) > 2:
        return category_path[2]
    return " / ".join(category_path)


def comparison_row(
    ours: dict | None,
    ours_alias: str | None,
    category: str | None,
    koch: dict | None,
    tolerance: float,
) -> dict:
    """One row of the sheet. `ours` is a product of decomposition.json, `koch` one of
    `read_koch`, either may be missing."""
    row = {
        "ours_name": ours and ours["product"],
        "koch_name": koch and koch["name"],
        "ours_alias": ours_alias,
        "category": category,
        "ours_unit": ours and ours["unit"],
    }
    ours_values = ours and {**ours, **main_raws(ours["raw"])}
    koch_values = koch and {**koch, **main_raws(koch_raw(koch))}
    for side, values in (("ours", ours_values), ("koch", koch_values)):
        row[f"{side}_main_raw_activity"] = values and values["main_raw_activity"]
        row[f"{side}_main_ingredient"] = values and values["main_ingredient"]
    delta_quantities = []
    for quantity in QUANTITIES:
        ours_value = ours_values[quantity] if ours_values else None
        koch_value = koch_values[quantity] if koch_values else None
        row[f"ours_{quantity}"] = significant(ours_value)
        row[f"koch_{quantity}"] = significant(koch_value)
        if ours_value is None or koch_value is None:
            continue
        delta = delta_pct(ours_value, koch_value)
        row[f"delta_pct_{quantity}"] = delta
        off = delta is None or abs(delta) > tolerance * 100
        if off and quantity in STATUS_QUANTITIES:
            delta_quantities.append(quantity)
    if ours:
        row["ours_main_ingredient_share"] = significant(ours["main_ingredient_share"])
        row["ours_raw_count"] = len(ours["raw"])
        row["ours_truncated"] = len(ours["truncated"])
    if row["ours_main_ingredient"] and row["koch_main_ingredient"]:
        same_main_ingredient = (
            row["ours_main_ingredient"] == row["koch_main_ingredient"]
        )
        row["same_main_ingredient"] = "yes" if same_main_ingredient else "no"
    row["delta_quantities"] = ", ".join(delta_quantities) or None
    if ours and koch:
        row["status"] = "delta" if delta_quantities else "match"
    else:
        row["status"] = "ours only" if ours else "koch only"
    return row


def compare(
    ours_products: dict[str, dict],
    koch_products: dict[str, dict],
    categories: dict[str, str],
    tolerance: float,
) -> list[dict]:
    """One row per product of either side, joined on the activity name, sorted by
    status then name. `ours_products` is keyed on the catalog alias, or on the activity
    name outside the catalog. `categories` gives the category of our products."""
    ours_names = set()
    rows = []
    for key, ours in ours_products.items():
        name = ours["product"]
        ours_names.add(name)
        rows.append(
            comparison_row(
                ours,
                key if key != name else None,
                categories.get(name),
                koch_products.get(name),
                tolerance,
            )
        )
    rows += [
        comparison_row(None, None, None, koch, tolerance)
        for name, koch in koch_products.items()
        if name not in ours_names
    ]
    return sorted(
        rows,
        key=lambda row: (
            STATUSES.index(row["status"]),
            row["ours_name"] or row["koch_name"],
        ),
    )


def ours_inputs(ours: dict) -> dict[tuple[str, str], float]:
    """Our inputs by (class, name). Waste goes with other, as Koch puts wastewater there."""
    inputs = {
        (input_class, name): amount
        for input_class, key in DETAILED_CLASSES.items()
        for name, amount in ours[key].items()
    }
    for name, amount in {**ours["other"], **ours["waste"]}.items():
        inputs["other", name] = amount
    return inputs


def input_rows(ours: dict | None, koch: dict) -> list[dict]:
    """The inputs of one product on both sides, aligned on (class, name), sorted by class
    then largest amount. Water, biowaste and transport are one total row per side.
    `ours` is None for a product on Koch's side only."""
    ours_amounts = ours_inputs(ours) if ours else {}
    koch_amounts = {
        key: amount
        for key, amount in koch["inputs"].items()
        if key[0] not in TOTAL_ONLY_CLASSES
    }
    for input_class, key in TOTAL_ONLY_CLASSES.items():
        koch_total = sum(
            amount
            for (koch_class, _), amount in koch["inputs"].items()
            if koch_class == input_class
        )
        ours_total = ours[key] if ours else None
        if ours_total or koch_total:
            if ours_total is not None:
                ours_amounts[input_class, TOTAL_INPUT_NAME] = ours_total
            koch_amounts[input_class, TOTAL_INPUT_NAME] = koch_total

    rows = []
    for input_class, name in set(ours_amounts) | set(koch_amounts):
        ours_amount = ours_amounts.get((input_class, name))
        koch_amount = koch_amounts.get((input_class, name))
        if ours_amount is not None and koch_amount is not None:
            side, delta = "both", delta_pct(ours_amount, koch_amount)
        else:
            side, delta = ("ours only" if koch_amount is None else "koch only"), None
        rows.append(
            {
                "ours_name": ours and ours["product"],
                "koch_name": koch["name"],
                "class": input_class,
                "input_name": name,
                "ours_amount": significant(ours_amount),
                "koch_amount": significant(koch_amount),
                "delta_pct": delta,
                "side": side,
            }
        )
    return sorted(
        rows,
        key=lambda row: (
            CLASS_ORDER.index(row["class"]),
            -max(row["ours_amount"] or 0, row["koch_amount"] or 0),
            row["input_name"],
        ),
    )


def write_sheet(report, title: str, columns: list[str], rows: list[dict]) -> None:
    sheet = report.create_sheet(title)
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions


def main(
    koch_workbook: Path = KOCH_WORKBOOK,
    ours_json: Path = OUTPUT_DIR / "decomposition.json",
    out: Path = OUTPUT_DIR / "ours_vs_koch.xlsx",
    tolerance: float = 0.01,
):
    """Compare decomposition.json (ours) with the ecolysis workbook (koch)."""
    workbook = openpyxl.load_workbook(koch_workbook, read_only=True, data_only=True)
    koch_products = read_koch(
        workbook[TOTALS_SHEET].iter_rows(min_row=TOTALS_FIRST_ROW, values_only=True),
        workbook[DETAIL_SHEET].iter_rows(min_row=DETAIL_FIRST_ROW, values_only=True),
    )
    ours_products = json.loads(ours_json.read_text())
    bw2data.projects.set_current(settings.bw.project)
    graph = BrightwayGraph(bw2data.Database(settings.bw.agribalyse))
    categories = {}
    for ours in ours_products.values():
        node = graph.node(ours["product"])
        if node:
            categories[ours["product"]] = category_name(node.category_path)
    rows = compare(ours_products, koch_products, categories, tolerance)

    # Detail for every Koch product, in the order of `rows` (delta first)
    ours_by_name = {product["product"]: product for product in ours_products.values()}
    inputs = []
    for row in rows:
        if row["koch_name"] is None:
            continue
        ours = ours_by_name.get(row["ours_name"])
        koch = koch_products[row["koch_name"]]
        inputs += input_rows(ours, koch)

    report = openpyxl.Workbook()
    report.remove(report.active)
    write_sheet(report, "ours_vs_koch", COLUMNS, rows)
    write_sheet(report, "inputs", INPUT_COLUMNS, inputs)
    report.save(out)
    counts = ", ".join(
        f"{status} {sum(row['status'] == status for row in rows)}"
        for status in STATUSES
    )
    print(f"{out}: {counts}, {len(inputs)} input rows")


if __name__ == "__main__":
    typer.run(main)
